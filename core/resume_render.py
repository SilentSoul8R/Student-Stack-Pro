"""Resume rendering: HTML preview, PDF (reportlab) and DOCX (python-docx) from one data model."""
from __future__ import annotations

import html
import io
from xml.sax.saxutils import escape as xesc


def lines(v) -> list[str]:
    if isinstance(v, list):
        v = "\n".join(str(x) for x in v)
    out = []
    for ln in str(v or "").splitlines():
        ln = ln.strip().lstrip("-•*·– ").strip()
        if ln:
            out.append(ln)
    return out


def prepare(r: dict) -> dict:
    g = lambda k: str(r.get(k, "") or "").strip()  # noqa: E731
    exp = [dict(role=str(e.get("role", "")).strip(), company=str(e.get("company", "")).strip(),
                location=str(e.get("location", "")).strip(),
                dates=" – ".join(x for x in (str(e.get("start", "")).strip(), str(e.get("end", "")).strip()) if x),
                bullets=lines(e.get("bullets"))) for e in r.get("experience", [])
           if str(e.get("role", "")).strip() or str(e.get("company", "")).strip()]
    edu = [dict(degree=str(e.get("degree", "")).strip(), school=str(e.get("school", "")).strip(),
                location=str(e.get("location", "")).strip(), year=str(e.get("year", "")).strip(),
                details=lines(e.get("details"))) for e in r.get("education", [])
           if str(e.get("degree", "")).strip() or str(e.get("school", "")).strip()]
    proj = [dict(name=str(p.get("name", "")).strip(), link=str(p.get("link", "")).strip(),
                 bullets=lines(p.get("bullets"))) for p in r.get("projects", []) if str(p.get("name", "")).strip()]
    contact = [x for x in (g("email"), g("phone"), g("location")) if x] + lines(r.get("links"))
    return dict(name=g("name"), title=g("title"), contact=contact, summary=g("summary"), experience=exp,
                education=edu, projects=proj, skills=lines(r.get("skills")), certs=lines(r.get("certifications")))


# ------------------------------------------------------------------ HTML
def to_html(r: dict, template: str = "modern", accent: str = "#4f46e5", compact: bool = False) -> str:
    d = prepare(r)
    e = html.escape
    modern = template == "modern"
    font = "'Segoe UI',Helvetica,Arial,sans-serif" if modern else "Georgia,'Times New Roman',serif"
    pad = "14px" if compact else "20px"
    css = f"""
body{{margin:0;background:#e5e7eb;font-family:{font};color:#222;font-size:{'12.5px' if compact else '13.5px'};line-height:1.45}}
.page{{max-width:794px;margin:12px auto;background:#fff;box-shadow:0 4px 18px rgba(0,0,0,.15);min-height:1000px}}
.band{{background:{accent};color:#fff;padding:28px 36px}}
.band h1{{margin:0;font-size:30px;letter-spacing:.5px}}.band .t{{font-size:15px;opacity:.95;margin-top:2px}}
.band .c{{font-size:11.5px;opacity:.9;margin-top:8px}}
.head{{text-align:center;padding:30px 36px 6px}}.head h1{{margin:0;font-size:30px;letter-spacing:1px}}
.head .t{{font-size:15px;color:{accent};margin-top:2px}}.head .c{{font-size:11.5px;color:#555;margin-top:6px}}
.body{{padding:8px 36px 30px}}
h2{{font-size:12.5px;letter-spacing:1.6px;text-transform:uppercase;color:{accent};border-bottom:1.5px solid {accent};
padding-bottom:3px;margin:{pad} 0 8px}}
.row{{display:flex;justify-content:space-between;gap:12px;margin-top:8px}}.row b{{font-size:1.02em}}
.dt{{color:#666;font-size:.9em;white-space:nowrap}}.sub{{color:#555;font-style:italic}}
ul{{margin:3px 0 0 18px;padding:0}}li{{margin:1px 0}}.chip{{display:inline-block;background:{accent}18;color:{accent};
border-radius:12px;padding:2px 10px;margin:2px 3px 2px 0;font-size:.92em}}
"""
    contact = " &nbsp;|&nbsp; ".join(e(c) for c in d["contact"])
    head = (f'<div class="band"><h1>{e(d["name"] or "Your Name")}</h1><div class="t">{e(d["title"])}</div>'
            f'<div class="c">{contact}</div></div>' if modern else
            f'<div class="head"><h1>{e(d["name"] or "Your Name")}</h1><div class="t">{e(d["title"])}</div>'
            f'<div class="c">{contact}</div></div>')
    body = []
    if d["summary"]:
        body.append(f"<h2>Profile</h2><div>{e(d['summary'])}</div>")
    if d["experience"]:
        body.append("<h2>Experience</h2>")
        for x in d["experience"]:
            sub = " · ".join(v for v in (x["company"], x["location"]) if v)
            body.append(f'<div class="row"><div><b>{e(x["role"])}</b>'
                        f'{" — " + e(sub) if sub else ""}</div><div class="dt">{e(x["dates"])}</div></div>')
            if x["bullets"]:
                body.append("<ul>" + "".join(f"<li>{e(b)}</li>" for b in x["bullets"]) + "</ul>")
    if d["projects"]:
        body.append("<h2>Projects</h2>")
        for p in d["projects"]:
            body.append(f'<div class="row"><div><b>{e(p["name"])}</b></div><div class="dt">{e(p["link"])}</div></div>')
            if p["bullets"]:
                body.append("<ul>" + "".join(f"<li>{e(b)}</li>" for b in p["bullets"]) + "</ul>")
    if d["education"]:
        body.append("<h2>Education</h2>")
        for x in d["education"]:
            sub = " · ".join(v for v in (x["school"], x["location"]) if v)
            body.append(f'<div class="row"><div><b>{e(x["degree"])}</b>{" — " + e(sub) if sub else ""}</div>'
                        f'<div class="dt">{e(x["year"])}</div></div>')
            if x["details"]:
                body.append("<ul>" + "".join(f"<li>{e(b)}</li>" for b in x["details"]) + "</ul>")
    if d["skills"]:
        body.append("<h2>Skills</h2>")
        for s in d["skills"]:
            if ":" in s:
                cat, items = s.split(":", 1)
                body.append(f"<div style='margin:3px 0'><b>{e(cat)}:</b> "
                            + "".join(f'<span class="chip">{e(i.strip())}</span>' for i in items.split(",") if i.strip())
                            + "</div>")
            else:
                body.append("".join(f'<span class="chip">{e(i.strip())}</span>' for i in s.split(",") if i.strip()))
    if d["certs"]:
        body.append("<h2>Certifications</h2><ul>" + "".join(f"<li>{e(c)}</li>" for c in d["certs"]) + "</ul>")
    return f"<!doctype html><html><head><meta charset='utf-8'><style>{css}</style></head><body><div class='page'>" \
           f"{head}<div class='body'>{''.join(body)}</div></div></body></html>"


# ------------------------------------------------------------------ PDF
def to_pdf(r: dict, template: str = "modern", accent: str = "#4f46e5") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.utils import simpleSplit
    from reportlab.platypus import (BaseDocTemplate, Frame, HRFlowable, KeepTogether, NextPageTemplate,
                                    PageTemplate, Paragraph, Spacer, Table, TableStyle)

    d = prepare(r)
    modern = template == "modern"
    F, FB, FI = ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique") if modern else (
        "Times-Roman", "Times-Bold", "Times-Italic")
    acc = colors.HexColor(accent)
    W, H = A4
    M = 16 * mm
    AW = W - 2 * M
    contact_txt = "   |   ".join(d["contact"])
    clines = simpleSplit(contact_txt, F, 9, AW) if contact_txt else []
    band_h = 22 + 28 + (18 if d["title"] else 0) + 12 * len(clines) + 18

    def draw_band(c, doc):
        c.saveState()
        c.setFillColor(acc)
        c.rect(0, H - band_h, W, band_h, fill=1, stroke=0)
        c.setFillColor(colors.white)
        y = H - 22 - 22
        c.setFont(FB, 24)
        c.drawString(M, y, d["name"] or "Your Name")
        if d["title"]:
            y -= 20
            c.setFont(F, 12)
            c.drawString(M, y, d["title"])
        c.setFont(F, 9)
        c.setFillColor(colors.HexColor("#e0e7ff"))
        for ln in clines:
            y -= 13
            c.drawString(M, y, ln)
        c.restoreState()

    body = ParagraphStyle("body", fontName=F, fontSize=9.5, leading=13, textColor=colors.HexColor("#222222"))
    small = ParagraphStyle("small", parent=body, fontSize=9, textColor=colors.HexColor("#666666"), alignment=TA_RIGHT)
    h = ParagraphStyle("h", parent=body, fontName=FB, fontSize=10.5, textColor=acc, spaceBefore=10, spaceAfter=1)
    bl = ParagraphStyle("bl", parent=body, leftIndent=12, bulletIndent=2, spaceBefore=1)

    def section(title):
        return [Paragraph(xesc(title.upper()), h),
                HRFlowable(width="100%", thickness=0.8, color=acc, spaceBefore=1, spaceAfter=4)]

    def head_row(left, right):
        t = Table([[Paragraph(left, body), Paragraph(xesc(right), small)]], colWidths=[AW * 0.74, AW * 0.26])
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                               ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 5),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
        return t

    def bullets(items):
        return [Paragraph(xesc(b), bl, bulletText="•") for b in items]

    story: list = []
    if modern:
        story.append(NextPageTemplate("later"))
    else:
        cs = ParagraphStyle("c", parent=body, alignment=TA_CENTER)
        story += [Paragraph(f"<font name='{FB}' size='24'>{xesc(d['name'] or 'Your Name')}</font>", cs)]
        if d["title"]:
            story.append(Paragraph(f"<font color='{accent}' size='12'>{xesc(d['title'])}</font>", cs))
        if contact_txt:
            story.append(Paragraph(f"<font size='9' color='#555555'>{xesc(contact_txt)}</font>", cs))
        story.append(Spacer(1, 4))
    if d["summary"]:
        story += section("Profile") + [Paragraph(xesc(d["summary"]), body)]
    if d["experience"]:
        story += section("Experience")
        for x in d["experience"]:
            sub = " · ".join(v for v in (x["company"], x["location"]) if v)
            left = f"<font name='{FB}'>{xesc(x['role'])}</font>" + (f" — {xesc(sub)}" if sub else "")
            bl_items = bullets(x["bullets"])
            story.append(KeepTogether([head_row(left, x["dates"])] + bl_items[:1]))
            story += bl_items[1:]
    if d["projects"]:
        story += section("Projects")
        for p in d["projects"]:
            bl_items = bullets(p["bullets"])
            story.append(KeepTogether([head_row(f"<font name='{FB}'>{xesc(p['name'])}</font>", p["link"])] + bl_items[:1]))
            story += bl_items[1:]
    if d["education"]:
        story += section("Education")
        for x in d["education"]:
            sub = " · ".join(v for v in (x["school"], x["location"]) if v)
            left = f"<font name='{FB}'>{xesc(x['degree'])}</font>" + (f" — {xesc(sub)}" if sub else "")
            bl_items = bullets(x["details"])
            story.append(KeepTogether([head_row(left, x["year"])] + bl_items[:1]))
            story += bl_items[1:]
    if d["skills"]:
        story += section("Skills")
        for s in d["skills"]:
            if ":" in s:
                cat, items = s.split(":", 1)
                story.append(Paragraph(f"<font name='{FB}'>{xesc(cat)}:</font> {xesc(items.strip())}", body))
            else:
                story.append(Paragraph(xesc(s), body))
    if d["certs"]:
        story += section("Certifications") + bullets(d["certs"])

    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=A4, leftMargin=M, rightMargin=M, topMargin=M, bottomMargin=M,
                          title=f"{d['name']} - Resume", author=d["name"])
    kw = dict(leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    if modern:
        first = Frame(M, M, AW, H - band_h - M - 6, id="f1", **kw)
        later = Frame(M, M, AW, H - 2 * M, id="f2", **kw)
        doc.addPageTemplates([PageTemplate(id="first", frames=[first], onPage=draw_band),
                              PageTemplate(id="later", frames=[later])])
    else:
        doc.addPageTemplates([PageTemplate(id="p", frames=[Frame(M, M, AW, H - 2 * M, id="f", **kw)])])
    doc.build(story)
    return buf.getvalue()


# ------------------------------------------------------------------ DOCX
def to_docx(r: dict, accent: str = "#4f46e5") -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    d = prepare(r)
    rgb = RGBColor.from_string(accent.lstrip("#").upper())
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    def heading(text):
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold, run.font.size, run.font.color.rgb = True, Pt(11), rgb
        return p

    p = doc.add_paragraph()
    run = p.add_run(d["name"] or "Your Name")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(24), rgb
    if d["title"]:
        doc.add_paragraph(d["title"])
    if d["contact"]:
        doc.add_paragraph("  |  ".join(d["contact"]))
    if d["summary"]:
        heading("Profile")
        doc.add_paragraph(d["summary"])
    if d["experience"]:
        heading("Experience")
        for x in d["experience"]:
            p = doc.add_paragraph()
            p.add_run(x["role"]).bold = True
            sub = " · ".join(v for v in (x["company"], x["location"]) if v)
            if sub:
                p.add_run(" — " + sub)
            if x["dates"]:
                p.add_run(f"   ({x['dates']})").italic = True
            for b in x["bullets"]:
                doc.add_paragraph(b, style="List Bullet")
    if d["projects"]:
        heading("Projects")
        for x in d["projects"]:
            p = doc.add_paragraph()
            p.add_run(x["name"]).bold = True
            if x["link"]:
                p.add_run(f"  {x['link']}")
            for b in x["bullets"]:
                doc.add_paragraph(b, style="List Bullet")
    if d["education"]:
        heading("Education")
        for x in d["education"]:
            p = doc.add_paragraph()
            p.add_run(x["degree"]).bold = True
            sub = " · ".join(v for v in (x["school"], x["location"]) if v)
            p.add_run((" — " + sub if sub else "") + (f"   ({x['year']})" if x["year"] else ""))
            for b in x["details"]:
                doc.add_paragraph(b, style="List Bullet")
    if d["skills"]:
        heading("Skills")
        for s in d["skills"]:
            doc.add_paragraph(s)
    if d["certs"]:
        heading("Certifications")
        for c in d["certs"]:
            doc.add_paragraph(c, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
