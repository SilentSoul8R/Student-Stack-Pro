"""Pure helpers: file readers, chunking, TF-IDF retrieval, markdown->docx, web snippets."""
from __future__ import annotations

import io
import re


# ------------------------------------------------------------------ readers
def read_pdf(data: bytes) -> list[tuple[int, str]]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:  # noqa: BLE001
            ok = 0
        if not ok:
            raise ValueError("This PDF is password-protected.")
    pages = []
    for i, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            text = ""
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        pages.append((i, text))
    return pages


def read_pptx(data: bytes) -> list[dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def walk(shapes, out):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes, out)
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                t = "\n".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())
                if t:
                    out.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))

    prs = Presentation(io.BytesIO(data))
    slides = []
    for n, slide in enumerate(prs.slides, 1):
        parts: list[str] = []
        walk(slide.shapes, parts)
        title = ""
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text.strip()
        except Exception:  # noqa: BLE001
            pass
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:  # noqa: BLE001
            pass
        slides.append(dict(n=n, title=title or f"Slide {n}", text="\n".join(parts), notes=notes))
    return slides


# ------------------------------------------------------------------ chunking + retrieval
def chunk_pages(pages: list[tuple[int, str]], size: int = 1100, overlap: int = 150, source: str = "") -> list[dict]:
    chunks = []
    step = max(size - overlap, 100)
    for page, text in pages:
        for start in range(0, len(text), step):
            piece = text[start:start + size].strip()
            if len(piece) >= 40:
                chunks.append(dict(page=page, text=piece, source=source))
    return chunks


class Retriever:
    """TF-IDF retriever: no GPU, no vector DB, no extra services."""

    def __init__(self, chunks: list[dict]):
        from sklearn.feature_extraction.text import TfidfVectorizer
        self.chunks = chunks
        texts = [c["text"] for c in chunks]
        try:
            self.vec = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True, max_features=60000)
            self.mat = self.vec.fit_transform(texts)
        except ValueError:  # e.g. only stop-words
            self.vec = TfidfVectorizer(sublinear_tf=True)
            self.mat = self.vec.fit_transform(texts)

    def search(self, query: str, k: int = 5) -> list[dict]:
        from sklearn.metrics.pairwise import linear_kernel
        sims = linear_kernel(self.vec.transform([query]), self.mat).ravel()
        order = sims.argsort()[::-1][:k]
        return [dict(self.chunks[i], score=float(sims[i])) for i in order if sims[i] > 0]


def group_by_chars(texts: list[str], max_chars: int) -> list[list[int]]:
    """Group consecutive items so each group's total length <= max_chars (items never split)."""
    groups: list[list[int]] = []
    cur: list[int] = []
    size = 0
    for i, t in enumerate(texts):
        if cur and size + len(t) > max_chars:
            groups.append(cur)
            cur, size = [], 0
        cur.append(i)
        size += len(t)
    if cur:
        groups.append(cur)
    return groups


# ------------------------------------------------------------------ markdown -> docx
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def _runs(par, text: str) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            par.add_run(part[1:-1]).font.name = "Consolas"
        else:
            par.add_run(part)


def md_to_docx(md: str, title: str = "") -> bytes:
    from docx import Document
    doc = Document()
    if title:
        doc.add_heading(title, 0)
    for line in md.splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        if s.strip() in ("---", "***"):
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            doc.add_heading(m.group(2).strip(), min(len(m.group(1)), 3))
        elif re.match(r"^\s*[-*•]\s+", s):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", s))
        elif re.match(r"^\s*\d+[.)]\s+", s):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", s))
        elif s.lstrip().startswith(">"):
            _runs(doc.add_paragraph(style="Intense Quote"), s.lstrip()[1:].strip())
        else:
            _runs(doc.add_paragraph(), s.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ optional web search
def web_available() -> bool:
    for mod in ("ddgs", "duckduckgo_search"):
        try:
            __import__(mod)
            return True
        except Exception:  # noqa: BLE001
            continue
    return False


def web_snippets(query: str, n: int = 6) -> list[tuple[str, str, str]]:
    try:
        try:
            from ddgs import DDGS
        except ImportError:
            from duckduckgo_search import DDGS
        with DDGS() as d:
            res = list(d.text(query, max_results=n))
        return [(r.get("title", ""), r.get("href") or r.get("url", ""), (r.get("body") or "")[:400]) for r in res]
    except Exception:  # noqa: BLE001
        return []


# ------------------------------------------------------------------ markdown -> PDF
def _pdf_safe(s: str) -> str:
    """Built-in PDF fonts only cover Latin-1/cp1252; replace anything else so it never crashes."""
    s = s.replace("\u2011", "-").replace("\u00a0", " ").replace("\t", "    ")
    return s.encode("cp1252", "replace").decode("cp1252")


def _inline_pdf(text: str) -> str:
    from xml.sax.saxutils import escape
    t = escape(_pdf_safe(text))
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<![\*\w])\*(?!\s)(.+?)(?<!\s)\*(?![\*\w])", r"<i>\1</i>", t)
    t = re.sub(r"`(.+?)`", r'<font name="Courier">\1</font>', t)
    return t


def md_to_pdf(md: str, title: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    ink = colors.HexColor("#1f2937")
    acc = colors.HexColor("#4f46e5")
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=14.5, textColor=ink, spaceAfter=4)
    hs = {1: ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=acc,
                            spaceBefore=12, spaceAfter=6),
          2: ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=14, leading=18, textColor=acc,
                            spaceBefore=10, spaceAfter=4),
          3: ParagraphStyle("h3", parent=body, fontName="Helvetica-Bold", fontSize=11.5, leading=15, spaceBefore=8,
                            spaceAfter=3)}
    quote = ParagraphStyle("q", parent=body, leftIndent=12, textColor=colors.HexColor("#4b5563"),
                           borderPadding=(2, 2, 2, 6), borderColor=acc, borderWidth=0, fontName="Helvetica-Oblique")
    code = ParagraphStyle("c", parent=body, fontName="Courier", fontSize=8.5, leading=11, backColor=colors.HexColor("#f3f4f6"),
                          leftIndent=6)

    def para(text, style, **kw):
        try:
            return Paragraph(_inline_pdf(text), style, **kw)
        except Exception:  # malformed inline markup -> plain text
            from xml.sax.saxutils import escape
            return Paragraph(escape(_pdf_safe(text)), style, **kw)

    story: list = []
    if title:
        story.append(Paragraph(_inline_pdf(title).replace("<b>", "").replace("</b>", ""),
                               ParagraphStyle("t", parent=hs[1], fontSize=24, leading=28, spaceAfter=10)))
        story.append(HRFlowable(width="100%", thickness=1.2, color=acc, spaceAfter=8))
    lines, i, in_code = md.splitlines(), 0, False
    while i < len(lines):
        s = lines[i].rstrip()
        i += 1
        if s.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            story.append(Paragraph(_inline_pdf(s).replace("<b>", "").replace("</b>", "") or "&nbsp;", code))
            continue
        if not s.strip():
            story.append(Spacer(1, 4))
            continue
        if s.strip() in ("---", "***", "___"):
            story.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#d1d5db"), spaceBefore=4, spaceAfter=6))
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            story.append(para(m.group(2).strip(), hs[min(len(m.group(1)), 3)]))
            continue
        if s.lstrip().startswith("|"):
            rows = [s]
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(lines[i].rstrip())
                i += 1
            data = []
            for r in rows:
                cells = [c.strip() for c in r.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    continue
                data.append([para(c, ParagraphStyle("tc", parent=body, fontSize=9, leading=12)) for c in cells])
            if data:
                width = max(len(r) for r in data)
                data = [r + [""] * (width - len(r)) for r in data]
                tbl = Table(data, hAlign="LEFT", repeatRows=1)
                tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d1d5db")),
                                         ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2ff")),
                                         ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                story += [tbl, Spacer(1, 6)]
            continue
        mb = re.match(r"^(\s*)[-*•]\s+(?:\[( |x|X)\]\s+)?(.*)", s)
        if mb:
            lvl = len(mb.group(1)) // 2
            mark = "☐" if mb.group(2) == " " else "☑" if mb.group(2) else "•"
            mark = mark if mb.group(2) is None else {"☐": "[ ]", "☑": "[x]"}[mark]
            story.append(para(mb.group(3), ParagraphStyle("li", parent=body, leftIndent=14 + 12 * lvl, bulletIndent=2 + 12 * lvl,
                                                          spaceAfter=2), bulletText=mark))
            continue
        mn = re.match(r"^(\s*)(\d+)[.)]\s+(.*)", s)
        if mn:
            lvl = len(mn.group(1)) // 2
            story.append(para(mn.group(3), ParagraphStyle("ol", parent=body, leftIndent=16 + 12 * lvl, bulletIndent=0 + 12 * lvl,
                                                          spaceAfter=2), bulletText=f"{mn.group(2)}."))
            continue
        if s.lstrip().startswith(">"):
            story.append(para(s.lstrip()[1:].strip(), quote))
            continue
        story.append(para(s.strip(), body))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#9ca3af"))
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.drawString(18 * mm, 10 * mm, "Student Stack Pro")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=_pdf_safe(title) or "Document")
    doc.build(story or [Spacer(1, 1)], onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def safe_pdf(md: str, title: str = "") -> bytes:
    """Never raises: returns b'' if the PDF could not be built (buttons then stay disabled)."""
    try:
        return md_to_pdf(md, title)
    except Exception:  # noqa: BLE001
        return b""


# ------------------------------------------------------------------ printable flashcards
def flashcards_pdf(cards: list[dict], title: str = "Flashcards") -> bytes:
    """Duplex-ready sheets: 8 cards/page; the back page is column-mirrored so fronts and backs line up."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfgen import canvas as cv

    W, H = A4
    M, HEAD = 12 * mm, 10 * mm
    cw, ch = (W - 2 * M) / 2, (H - 2 * M - HEAD) / 4
    buf = io.BytesIO()
    c = cv.Canvas(buf, pagesize=A4)
    c.setTitle(_pdf_safe(title))

    def draw_text(text, x, y, w, h, font="Helvetica"):
        text = _pdf_safe(text)
        for size in (16, 14, 12, 11, 10, 9, 8):
            lines = simpleSplit(text, font, size, w - 14 * mm)
            if len(lines) * size * 1.25 <= h - 16 * mm:
                break
        else:
            max_lines = int((h - 16 * mm) // (8 * 1.25))
            lines = lines[:max_lines]
            if lines:
                lines[-1] = lines[-1][:-3].rstrip() + "..."
        block = len(lines) * size * 1.25
        ty = y + h / 2 + block / 2 - size
        c.setFont(font, size)
        for ln in lines:
            c.drawCentredString(x + w / 2, ty, ln)
            ty -= size * 1.25

    def page(batch, start, back):
        c.setFont("Helvetica", 8)
        c.setFillColor(colors.HexColor("#6b7280"))
        c.drawString(M, H - M - 2, f"{_pdf_safe(title)} - {'ANSWERS (flip on long edge)' if back else 'QUESTIONS (print this side first)'}")
        for k, card in enumerate(batch):
            row, col = divmod(k, 2)
            if back:
                col = 1 - col
            x, y = M + col * cw, H - M - HEAD - (row + 1) * ch
            c.setFillColor(colors.HexColor("#ecfdf5" if back else "#eef2ff"))
            c.setStrokeColor(colors.HexColor("#9ca3af"))
            c.setDash(3, 3)
            c.roundRect(x + 2, y + 2, cw - 4, ch - 4, 8, fill=1, stroke=1)
            c.setDash()
            c.setFillColor(colors.HexColor("#6b7280"))
            c.setFont("Helvetica", 7)
            c.drawString(x + 8, y + ch - 12, f"{'A' if back else 'Q'}{start + k + 1}")
            c.setFillColor(colors.HexColor("#111827"))
            draw_text(card["a"] if back else card["q"], x, y, cw, ch, "Helvetica-Bold" if not back else "Helvetica")
        c.showPage()

    for start in range(0, len(cards), 8):
        batch = cards[start:start + 8]
        page(batch, start, False)
        page(batch, start, True)
    if not cards:
        c.showPage()
    c.save()
    return buf.getvalue()


def flashcards_csv(cards: list[dict]) -> str:
    """Two-column CSV (front,back) that Quizlet/Anki/Google Sheets import directly."""
    import csv
    buf = io.StringIO()
    w = csv.writer(buf)
    for cd in cards:
        w.writerow([cd["q"], cd["a"]])
    return buf.getvalue()
